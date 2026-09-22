"""Read the supplied XLSX without modifying it; emit a compact JS data snapshot.

Usage: bundled-python import-structure-source.py [source.xlsx]
The output goes to stdout so the caller can install it with apply_patch.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


source = Path(sys.argv[1] if len(sys.argv) > 1 else
              '/Users/admin/Desktop/15092026_структура для подготовки мока данных.xlsx')
workbook = load_workbook(source, read_only=True, data_only=True)
sheet = workbook.active
values = list(sheet.iter_rows(values_only=True))
strings = []
indices = {}


def encode(value):
    if value is None:
        return -1
    text = str(value)
    if text not in indices:
        indices[text] = len(strings)
        strings.append(text)
    return indices[text]


rows = [[encode(value) for value in row] for row in values[1:]]
snapshot = {
    'file': source.name,
    'date': '2026-09-15',
    'sheet': sheet.title,
    'columns': list(values[0]),
    'strings': strings,
    'rows': rows,
}
print('/* Local snapshot of the supplied workbook. Indices reference strings; -1 means an empty cell. */')
print('window.BPM_STRUCTURE_SOURCE = ' + json.dumps(snapshot, ensure_ascii=False, separators=(',', ':')) + ';')
workbook.close()
